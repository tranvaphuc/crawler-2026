from pathlib import Path

from openpyxl import load_workbook


SOURCE = Path("/Users/phuctran/Downloads/OBRANDING_KOMPA_0509.xlsx")
OUTPUT = Path("outputs/kompa_profile_work/ids.txt")

workbook = load_workbook(SOURCE, read_only=True, data_only=True)
sheet = workbook["Kompa"]
header = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
id_index = next(i for i, value in enumerate(header) if str(value).strip().casefold() == "id")

ids = set()
blank_rows = 0
for row in sheet.iter_rows(min_row=2, values_only=True):
    value = row[id_index]
    if value is None or not str(value).strip():
        blank_rows += 1
    else:
        ids.add(str(value).strip())

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
OUTPUT.write_text("\n".join(sorted(ids)) + "\n", encoding="utf-8")
print({
    "sheet_rows": sheet.max_row - 1,
    "unique_ids": len(ids),
    "blank_id_rows": blank_rows,
    "output": str(OUTPUT),
})
workbook.close()
