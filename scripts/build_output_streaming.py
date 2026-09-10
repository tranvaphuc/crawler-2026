import csv
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill


SOURCE = Path("/Users/phuctran/Desktop/check_xxxx.xlsx")
WORK_DIR = Path("outputs/phone_language_work")
OUTPUT_DIR = Path("outputs/01a07164-8064-7313-962f-0d73887aa809")
OUTPUT = OUTPUT_DIR / "check_xxxx_phone_language.xlsx"

DARK_BLUE = "1F4E78"
WHITE = "FFFFFF"
DARK_TEXT = "1F2937"
MUTED_TEXT = "5B6573"


def styled_cell(ws, value, *, bold=False, color=DARK_TEXT, size=10, italic=False, fill=None, align=None):
    cell = WriteOnlyCell(ws, value=value)
    cell.font = Font(name="Arial", bold=bold, italic=italic, color=color, size=size)
    if fill:
        cell.fill = PatternFill("solid", fgColor=fill)
    if align:
        cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)
    return cell


def literal_cell(ws, value):
    if not isinstance(value, str):
        return value
    cell = WriteOnlyCell(ws, value=value)
    cell.data_type = "s"
    return cell


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    source_wb = load_workbook(SOURCE, read_only=True, data_only=True)
    source_ws = source_wb["Data"]

    out_wb = Workbook(write_only=True)
    data_ws = out_wb.create_sheet("Data")
    data_ws.freeze_panes = "A2"
    data_ws.auto_filter.ref = f"A1:E{source_ws.max_row}"
    widths = {"A": 80, "B": 45, "C": 22, "D": 28, "E": 14}
    for col, width in widths.items():
        data_ws.column_dimensions[col].width = width

    data_ws.append([
        styled_cell(data_ws, "Content", bold=True, color=WHITE, fill=DARK_BLUE, align="center"),
        styled_cell(data_ws, "SiteName", bold=True, color=WHITE, fill=DARK_BLUE, align="center"),
        styled_cell(data_ws, "SiteId", bold=True, color=WHITE, fill=DARK_BLUE, align="center"),
        styled_cell(data_ws, "Số điện thoại", bold=True, color=WHITE, fill=DARK_BLUE, align="center"),
        styled_cell(data_ws, "Language", bold=True, color=WHITE, fill=DARK_BLUE, align="center"),
    ])

    with (WORK_DIR / "row_results.csv").open("r", encoding="utf-8", newline="") as rf:
        result_reader = csv.DictReader(rf)
        source_rows = source_ws.iter_rows(min_row=2, values_only=True)
        written = 0
        for source_row, result in zip(source_rows, result_reader):
            data_ws.append([
                literal_cell(data_ws, source_row[0]),
                literal_cell(data_ws, source_row[1]),
                literal_cell(data_ws, source_row[2]),
                literal_cell(data_ws, result["phone_number"]),
                literal_cell(data_ws, result["language"]),
            ])
            written += 1

    top_ws = out_wb.create_sheet("Top 500 Việt Nam")
    top_ws.sheet_view.showGridLines = False
    top_ws.freeze_panes = "A5"
    top_ws.column_dimensions["A"].width = 8
    top_ws.column_dimensions["B"].width = 22
    top_ws.column_dimensions["C"].width = 56
    for col in ("D", "E", "F"):
        top_ws.column_dimensions[col].width = 22
    top_ws.append([styled_cell(top_ws, "Top 500 site có nhiều số điện thoại trong content tiếng Việt", bold=True, size=16)])
    top_ws.append([styled_cell(
        top_ws,
        "Xếp hạng theo số điện thoại duy nhất; dùng lượt xuất hiện và số content tiếng Việt để phân hạng khi bằng nhau.",
        italic=True,
        color=MUTED_TEXT,
    )])
    top_ws.append([])
    top_ws.append([
        styled_cell(top_ws, "Hạng", bold=True, color=WHITE, fill=DARK_BLUE, align="center"),
        styled_cell(top_ws, "SiteId", bold=True, color=WHITE, fill=DARK_BLUE, align="center"),
        styled_cell(top_ws, "SiteName", bold=True, color=WHITE, fill=DARK_BLUE, align="center"),
        styled_cell(top_ws, "Số điện thoại duy nhất", bold=True, color=WHITE, fill=DARK_BLUE, align="center"),
        styled_cell(top_ws, "Lượt xuất hiện SĐT", bold=True, color=WHITE, fill=DARK_BLUE, align="center"),
        styled_cell(top_ws, "Số content tiếng Việt", bold=True, color=WHITE, fill=DARK_BLUE, align="center"),
    ])

    top_count = 0
    with (WORK_DIR / "top_sites.csv").open("r", encoding="utf-8", newline="") as tf:
        for row in csv.DictReader(tf):
            top_ws.append([
                int(row["Rank"]),
                literal_cell(top_ws, row["SiteId"]),
                literal_cell(top_ws, row["SiteName"]),
                int(row["UniquePhoneCount"]),
                int(row["PhoneOccurrences"]),
                int(row["VietnameseContentRows"]),
            ])
            top_count += 1
    top_ws.auto_filter.ref = f"A4:F{4 + top_count}"

    out_wb.save(OUTPUT)
    source_wb.close()
    print({"output": str(OUTPUT.resolve()), "data_rows": written, "top_rows": top_count})


if __name__ == "__main__":
    main()
