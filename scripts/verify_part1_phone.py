import sys
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


SOURCE = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("/Users/phuctran/Desktop/part1.xlsx")
OUTPUT = Path(sys.argv[2]) if len(sys.argv) > 2 else Path("outputs/01a07164-8064-7313-962f-0d73887aa809/part1_phone.xlsx")
EXPECTED_PHONE_ROWS = int(sys.argv[3]) if len(sys.argv) > 3 else 216337
EXPECTED_PHONE_COUNT = int(sys.argv[4]) if len(sys.argv) > 4 else 237017

source_wb = load_workbook(SOURCE, read_only=True, data_only=True)
output_wb = load_workbook(OUTPUT, read_only=True, data_only=False)
source_ws = source_wb["Data"]
output_ws = output_wb["Data"]
assert output_wb.sheetnames == ["Data"]

source_header = tuple(cell.value for cell in next(source_ws.iter_rows(min_row=1, max_row=1)))
content_index = next(i for i, value in enumerate(source_header) if str(value).strip().casefold() == "content")
expected_header = source_header[: content_index + 1] + ("Phone",) + source_header[content_index + 1 :]
source_max_row = source_ws.max_row
TARGETS = {2, 3, 100000, source_max_row // 2, source_max_row}

source_samples = {}
for row_number, row in enumerate(source_ws.iter_rows(values_only=True), 1):
    if row_number in TARGETS:
        source_samples[row_number] = row

output_samples = {}
formula_count = 0
rows_with_phone = 0
phone_count = 0
for row_number, row in enumerate(output_ws.iter_rows(), 1):
    if row_number == 1:
        assert tuple(cell.value for cell in row) == expected_header
        header_phone = row[content_index + 1]
    if row_number in TARGETS:
        output_samples[row_number] = tuple(cell.value for cell in row)
    for cell in row:
        formula_count += int(cell.data_type == "f")
    phone_cell = row[content_index + 1]
    if row_number > 1 and phone_cell.value:
        rows_with_phone += 1
        phone_count += len(str(phone_cell.value).split("; "))

output_row_count = row_number
output_column_count = len(row)

for row_number in TARGETS:
    original = source_samples[row_number]
    enriched = output_samples[row_number]
    normalize = lambda values: tuple("" if value is None else value for value in values)
    assert normalize(enriched[: content_index + 1]) == normalize(original[: content_index + 1])
    assert normalize(enriched[content_index + 2 :]) == normalize(original[content_index + 1 :])

assert formula_count == 0
assert output_row_count == source_max_row
assert output_column_count == len(source_header) + 1
assert rows_with_phone == EXPECTED_PHONE_ROWS
assert phone_count == EXPECTED_PHONE_COUNT

assert header_phone.font.bold
assert header_phone.border.bottom.style == "medium"
last_column = get_column_letter(len(expected_header))
with ZipFile(OUTPUT) as archive:
    sheet_xml = archive.read("xl/worksheets/sheet1.xml")
    assert b'topLeftCell="A2"' in sheet_xml
    assert f'<autoFilter ref="A1:{last_column}{source_max_row}"'.encode() in sheet_xml

print({
    "shape": [output_row_count, output_column_count],
    "rows_with_phone": rows_with_phone,
    "phone_count": phone_count,
    "formula_count": formula_count,
    "verified_rows": sorted(TARGETS),
    "first_examples": [output_samples[n][content_index:content_index + 3] for n in sorted(TARGETS)[:2]],
})

source_wb.close()
output_wb.close()
