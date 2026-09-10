import csv
from pathlib import Path

from openpyxl import load_workbook


OUTPUT = Path("outputs/01a07164-8064-7313-962f-0d73887aa809/check_xxxx_phone_language.xlsx")
RESULTS = Path("outputs/phone_language_work/row_results.csv")


wb = load_workbook(OUTPUT, read_only=True, data_only=False)
assert wb.sheetnames == ["Data", "Top 500 Việt Nam"]
data = wb["Data"]
top = wb["Top 500 Việt Nam"]
data.calculate_dimension(force=True)
top.calculate_dimension(force=True)
assert data.max_row == 663351 and data.max_column == 5
assert top.max_row == 504 and top.max_column == 6

expected = {}
targets = {2, 3, 100000, 300000, 663351}
with RESULTS.open("r", encoding="utf-8", newline="") as f:
    for row in csv.DictReader(f):
        n = int(row["row_number"])
        if n in targets:
            expected[n] = (row["phone_number"], row["language"])

actual = {}
for n, row in enumerate(data.iter_rows(values_only=True), 1):
    if n == 1:
        assert row == ("Content", "SiteName", "SiteId", "Số điện thoại", "Language")
    if n in targets:
        actual[n] = (row[3] or "", row[4])
assert actual == expected

top_rows = list(top.iter_rows(min_row=4, max_row=8, values_only=True))
assert top_rows[0] == (
    "Hạng", "SiteId", "SiteName", "Số điện thoại duy nhất", "Lượt xuất hiện SĐT", "Số content tiếng Việt"
)
assert top_rows[1][0] == 1 and top_rows[1][3] >= top_rows[2][3]

formula_count = 0
for ws in (data, top):
    for row in ws.iter_rows():
        for cell in row:
            if cell.data_type == "f":
                formula_count += 1
assert formula_count == 0

print({
    "sheets": wb.sheetnames,
    "data_shape": [data.max_row, data.max_column],
    "top_shape": [top.max_row, top.max_column],
    "sample_rows_verified": sorted(targets),
    "formula_errors": 0,
    "top_preview": top_rows,
})
wb.close()
