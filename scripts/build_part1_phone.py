import re
import sys
import unicodedata
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


SOURCE = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("/Users/phuctran/Desktop/part1.xlsx")
OUTPUT_DIR = Path("outputs/01a07164-8064-7313-962f-0d73887aa809")
OUTPUT = OUTPUT_DIR / (sys.argv[2] if len(sys.argv) > 2 else "part1_phone.xlsx")

PHONE_CANDIDATE = re.compile(r"(?<!\d)(?:\+?\d(?:[\s()./_\-]*\d){6,14})(?!\d)")
PHONE_CUE = re.compile(
    r"(?:đt|điện\s*thoại|sđt|phone|tel|hotline|zalo|whatsapp|liên\s*hệ|lh|call|☎|📞|📱)",
    re.IGNORECASE,
)


def extract_phones(value):
    if value is None:
        return []
    text = unicodedata.normalize("NFKC", str(value))
    text = text.replace("\ufe0f", "").replace("\u20e3", "")
    phones = []
    seen = set()
    for match in PHONE_CANDIDATE.finditer(text):
        raw = match.group(0).strip()
        digits = re.sub(r"\D", "", raw)
        if not (7 <= len(digits) <= 15):
            continue
        groups = re.findall(r"\d+", raw)
        # Loại các giá tiền được nhóm theo hàng nghìn, ví dụ 17.500.000.
        if len(groups) >= 3 and all(len(group) == 3 for group in groups[1:]) and len(groups[0]) <= 3:
            if not raw.lstrip().startswith(("+", "0")):
                continue
        if re.fullmatch(r"(?:19|20)\d{2}[\s./_-]+\d{1,2}[\s./_-]+\d{1,2}", raw):
            continue

        starts_plus = raw.lstrip().startswith("+")
        nearby = text[max(0, match.start() - 35): min(len(text), match.end() + 35)]
        plausible = (
            (starts_plus and 10 <= len(digits) <= 15)
            or (digits.startswith("0") and 9 <= len(digits) <= 11)
            or (digits.startswith("84") and 10 <= len(digits) <= 12)
            or (8 <= len(digits) <= 12 and PHONE_CUE.search(nearby))
        )
        if not plausible:
            continue

        normalized = ("+" if starts_plus else "") + digits
        if normalized not in seen:
            seen.add(normalized)
            phones.append(normalized)
    return phones


def literal_cell(ws, value):
    if not isinstance(value, str):
        return value
    cell = WriteOnlyCell(ws, value=value)
    cell.data_type = "s"
    return cell


def header_cell(ws, value):
    cell = WriteOnlyCell(ws, value=value)
    cell.font = Font(name="Arial", size=11, bold=True, color="000000")
    cell.fill = PatternFill("solid", fgColor="FFFFFF")
    cell.alignment = Alignment(horizontal="left", vertical="center")
    cell.border = Border(
        left=Side(style="thin", color="D1D5DB"),
        right=Side(style="thin", color="D1D5DB"),
        top=Side(style="thin", color="D1D5DB"),
        bottom=Side(style="medium", color="10B981"),
    )
    return cell


def main():
    source_wb = load_workbook(SOURCE, read_only=True, data_only=True)
    source_ws = source_wb["Data"]
    source_rows = source_ws.iter_rows(values_only=True)
    headers = list(next(source_rows))
    content_index = next(i for i, header in enumerate(headers) if str(header).strip().casefold() == "content")
    output_headers = headers[: content_index + 1] + ["Phone"] + headers[content_index + 1 :]

    out_wb = Workbook(write_only=True)
    out_ws = out_wb.create_sheet("Data")
    out_ws.freeze_panes = "A2"
    out_ws.sheet_view.showGridLines = True
    out_ws.append([header_cell(out_ws, value) for value in output_headers])
    out_ws.row_dimensions[1].height = 24

    # Widths follow the compact table layout in the supplied screenshot.
    width_by_header = {
        "Id": 28, "Title": 55, "Content": 85, "Phone": 24, "Description": 38,
        "UrlComment": 42, "PublishedDate": 20, "SiteName": 42, "SiteId": 22,
        "Channel": 16, "UrlTopic": 42, "ParentId": 28, "Type": 18,
    }
    for index, header in enumerate(output_headers, 1):
        out_ws.column_dimensions[get_column_letter(index)].width = width_by_header.get(str(header), 22)

    total_rows = 0
    rows_with_phone = 0
    phone_count = 0
    for row in source_rows:
        phones = extract_phones(row[content_index])
        output_row = list(row[: content_index + 1]) + ["; ".join(phones)] + list(row[content_index + 1 :])
        out_ws.append([literal_cell(out_ws, value) for value in output_row])
        total_rows += 1
        rows_with_phone += int(bool(phones))
        phone_count += len(phones)

    last_column = get_column_letter(len(output_headers))
    out_ws.auto_filter.ref = f"A1:{last_column}{total_rows + 1}"
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    out_wb.save(OUTPUT)
    source_wb.close()
    print({
        "output": str(OUTPUT.resolve()),
        "rows": total_rows,
        "rows_with_phone": rows_with_phone,
        "phone_count": phone_count,
        "columns": len(output_headers),
    })


if __name__ == "__main__":
    main()
