import json
from pathlib import Path

from openpyxl import load_workbook


source = Path("/Users/phuctran/Downloads/Dashboard_28761_Top-Sources_2026-09-02-16-52-46.xlsx")
output = Path("outputs/top_sources_links.json")
workbook = load_workbook(source, read_only=False, data_only=False)
sheet = workbook["Top Sources"]
links = []
for row in range(2, sheet.max_row + 1):
    cell = sheet.cell(row, 2)
    links.append(cell.hyperlink.target if cell.hyperlink else "")
output.parent.mkdir(parents=True, exist_ok=True)
output.write_text(json.dumps(links, ensure_ascii=False), encoding="utf-8")
print({"rows": len(links), "links": sum(bool(link) for link in links), "output": str(output)})
