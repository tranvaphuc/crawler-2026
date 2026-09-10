from pathlib import Path
from openpyxl import load_workbook


SOURCE = Path("/Users/phuctran/Desktop/check_xxxx.xlsx")


wb = load_workbook(SOURCE, read_only=True, data_only=False)
print("sheets", wb.sheetnames)
for ws in wb.worksheets:
    print("sheet", ws.title, "rows", ws.max_row, "cols", ws.max_column)
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=min(ws.max_row, 8), values_only=True), 1):
        vals = [repr(v)[:300] for v in row]
        print(i, vals)
