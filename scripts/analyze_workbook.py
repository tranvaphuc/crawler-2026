import csv
import re
import unicodedata
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


SOURCE = Path("/Users/phuctran/Desktop/check_xxxx.xlsx")
WORK_DIR = Path("outputs/phone_language_work")
ROW_RESULTS = WORK_DIR / "row_results.csv"
TOP_RESULTS = WORK_DIR / "top_sites.csv"

VI_CORE_CHARS = set(
    "ăắằẳẵặđơớờởỡợưứừửữự"
    "ĂẮẰẲẴẶĐƠỚỜỞỠỢƯỨỪỬỮỰ"
)
VI_EXTENDED_CHARS = set(
    "âấầẩẫậêếềểễệôốồổỗộ"
    "ÂẤẦẨẪẬÊẾỀỂỄỆÔỐỒỔỖỘ"
)
VI_STRONG = re.compile(
    r"\b(việt\s*nam|tuyển\s*dụng|liên\s*hệ|điện\s*thoại|không|được|"
    r"chúng\s*tôi|khách\s*hàng|công\s*ty|sản\s*phẩm|thành\s*phố|"
    r"hà\s*nội|hồ\s*chí\s*minh|tp\.?\s*hcm|sđt)\b",
    re.IGNORECASE,
)
VI_COMMON = re.compile(
    r"\b(và|của|cho|bạn|với|trong|một|những|các|có|là|đến|tại|"
    r"mua|bán|giá|nhà|hàng|ngày|tháng|năm|dịch\s*vụ|zalo)\b",
    re.IGNORECASE,
)
VI_ASCII = re.compile(
    r"\b(khong|duoc|chung\s*toi|lien\s*he|tuyen\s*dung|khach\s*hang|"
    r"cong\s*ty|san\s*pham|viet\s*nam|ha\s*noi|ho\s*chi\s*minh|"
    r"mua|ban|gia|nha|cho|cua|voi|zalo)\b",
    re.IGNORECASE,
)
PHONE_CANDIDATE = re.compile(r"(?<!\d)(?:\+?\d(?:[\s()./\-]*\d){6,14})(?!\d)")


def is_vietnamese(value):
    if value is None:
        return False
    text = unicodedata.normalize("NFC", str(value))
    if not text.strip():
        return False
    strong = len(VI_STRONG.findall(text))
    common = len(VI_COMMON.findall(text))
    ascii_hits = len(VI_ASCII.findall(text))
    core_chars = sum(ch in VI_CORE_CHARS for ch in text)
    extended_chars = sum(ch in VI_EXTENDED_CHARS for ch in text)
    letters = sum(ch.isalpha() for ch in text)
    if strong >= 1:
        return True
    if core_chars >= 1 and common >= 1:
        return True
    if core_chars >= 3:
        return True
    if extended_chars >= 2 and common >= 2:
        return True
    if letters >= 35 and ascii_hits >= 4:
        return True
    return False


def extract_phones(value):
    if value is None:
        return []
    text = unicodedata.normalize("NFKC", str(value))
    text = text.replace("\ufe0f", "").replace("\u20e3", "")
    found = []
    seen = set()
    for match in PHONE_CANDIDATE.finditer(text):
        raw = match.group(0).strip()
        digits = re.sub(r"\D", "", raw)
        if not (7 <= len(digits) <= 15):
            continue
        groups = re.findall(r"\d+", raw)
        # Exclude obvious grouped monetary values such as 17.500.000.
        if len(groups) >= 3 and all(len(g) == 3 for g in groups[1:]) and len(groups[0]) <= 3:
            if not raw.lstrip().startswith(("+", "0")):
                continue
        # Exclude strings that are predominantly a run of years/dates.
        if re.fullmatch(r"(?:19|20)\d{2}[\s./-]+\d{1,2}[\s./-]+\d{1,2}", raw):
            continue
        normalized = ("+" if raw.lstrip().startswith("+") else "") + digits
        if normalized not in seen:
            seen.add(normalized)
            found.append(normalized)
    return found


def main():
    WORK_DIR.mkdir(parents=True, exist_ok=True)
    wb = load_workbook(SOURCE, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    headers = [str(v).strip() if v is not None else "" for v in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    index = {name.casefold(): i for i, name in enumerate(headers)}
    content_i = index["content"]
    site_name_i = index["sitename"]
    site_id_i = index["siteid"]

    stats = defaultdict(lambda: {"phones": set(), "occurrences": 0, "vi_rows": 0})
    total_rows = 0
    vi_rows = 0
    phone_rows = 0

    with ROW_RESULTS.open("w", encoding="utf-8", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(["row_number", "phone_number", "language"])
        for row_number, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
            content = row[content_i]
            phones = extract_phones(content)
            vi = is_vietnamese(content)
            language = "Vietnamese" if vi else "Other"
            writer.writerow([row_number, "; ".join(phones), language])
            total_rows += 1
            vi_rows += int(vi)
            phone_rows += int(bool(phones))
            if vi and phones:
                site_id = "" if row[site_id_i] is None else str(row[site_id_i])
                site_name = "" if row[site_name_i] is None else str(row[site_name_i])
                key = (site_id, site_name)
                rec = stats[key]
                rec["phones"].update(phones)
                rec["occurrences"] += len(phones)
                rec["vi_rows"] += 1

    ranked = sorted(
        stats.items(),
        key=lambda kv: (-len(kv[1]["phones"]), -kv[1]["occurrences"], -kv[1]["vi_rows"], kv[0][0]),
    )[:500]
    with TOP_RESULTS.open("w", encoding="utf-8", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(["Rank", "SiteId", "SiteName", "UniquePhoneCount", "PhoneOccurrences", "VietnameseContentRows"])
        for rank, ((site_id, site_name), rec) in enumerate(ranked, 1):
            writer.writerow([rank, site_id, site_name, len(rec["phones"]), rec["occurrences"], rec["vi_rows"]])

    print({
        "rows": total_rows,
        "vietnamese_rows": vi_rows,
        "phone_rows": phone_rows,
        "sites_with_vi_phones": len(stats),
        "top_rows": len(ranked),
        "row_results": str(ROW_RESULTS),
        "top_results": str(TOP_RESULTS),
    })


if __name__ == "__main__":
    main()
