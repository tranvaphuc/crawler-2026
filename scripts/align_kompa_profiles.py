import json
from pathlib import Path

from openpyxl import load_workbook


SOURCE = Path("/Users/phuctran/Downloads/OBRANDING_KOMPA_0509.xlsx")
PROFILES = Path("outputs/kompa_profile_work/profiles.ndjson")
OUTPUT = Path("outputs/kompa_profile_work/profiles_aligned.ndjson")

profiles = {}
with PROFILES.open("r", encoding="utf-8") as file:
    for line in file:
        if line.strip():
            record = json.loads(line)
            profiles[record["id"]] = (record["authorId"], record["author"])

workbook = load_workbook(SOURCE, read_only=True, data_only=True)
sheet = workbook["Kompa"]
header = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
id_index = next(i for i, value in enumerate(header) if str(value).strip().casefold() == "id")
matched = 0
missing = 0
with OUTPUT.open("w", encoding="utf-8") as output:
    for row in sheet.iter_rows(min_row=2, values_only=True):
        doc_id = str(row[id_index]).strip()
        profile = profiles.get(doc_id)
        if profile is None:
            missing += 1
            profile = ("", "")
        else:
            matched += 1
        output.write(json.dumps({"authorId": profile[0], "author": profile[1]}, ensure_ascii=False) + "\n")

print({"rows": matched + missing, "matched": matched, "missing": missing, "output": str(OUTPUT)})
workbook.close()
